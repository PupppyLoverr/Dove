#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Create, inspect and edit .xlsx workbooks. See ../SKILL.md."""
import argparse
import csv
import datetime as dt
import json
import re
import sys

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, range_boundaries
    from openpyxl.workbook.properties import CalcProperties
except ImportError:
    sys.exit("openpyxl is missing: run with `uv run` or `python3 -m pip install --user openpyxl`")

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
HEADER_BORDER = Border(bottom=Side(style="thin", color="808080"))
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def coerce(v):
    if isinstance(v, str):
        if DATE.match(v):
            try:
                return dt.date.fromisoformat(v)
            except ValueError:
                return v
        return v
    return v


def parse_scalar(text):
    if text.startswith("="):
        return text
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        pass
    if text.lower() in ("true", "false"):
        return text.lower() == "true"
    return coerce(text)


def recalc_on_open(wb):
    wb.calculation = CalcProperties(fullCalcOnLoad=True)


def build_sheet(ws, spec):
    cols = spec.get("columns") or []
    r = 1
    if cols:
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=c.get("header", ""))
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, HEADER_BORDER
            cell.alignment = Alignment(vertical="center")
            if c.get("width"):
                ws.column_dimensions[get_column_letter(j)].width = c["width"]
        r = 2
    rows = spec.get("rows") or []
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r + i, column=j, value=coerce(v))
            fmt = cols[j - 1].get("format") if j - 1 < len(cols) else None
            if fmt:
                cell.number_format = fmt
            elif isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"
    last = r + len(rows) - 1
    totals = spec.get("totals")
    if totals and rows:
        tr = last + 1
        label = ws.cell(row=tr, column=1, value=totals.get("label", "Total"))
        label.font = HEADER_FONT
        for col in totals.get("sum", []):
            j = openpyxl.utils.column_index_from_string(col)
            cell = ws.cell(row=tr, column=j, value=f"=SUM({col}{r}:{col}{last})")
            cell.font = HEADER_FONT
            cell.border = Border(top=Side(style="thin", color="808080"))
            fmt = cols[j - 1].get("format") if j - 1 < len(cols) else None
            if fmt:
                cell.number_format = fmt
        last = tr
    if not cols:
        for j in range(1, ws.max_column + 1):
            letter = get_column_letter(j)
            width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
            ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
    if spec.get("freeze"):
        ws.freeze_panes = spec["freeze"]
    if spec.get("filter") and cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{r + len(rows) - 1}"
    for h in spec.get("highlight", []):
        fill = PatternFill("solid", fgColor=h.get("fill", "EDEDED"))
        if "above" in h:
            ws.conditional_formatting.add(h["range"], CellIsRule(operator="greaterThan", formula=[str(h["above"])], fill=fill))
        if "below" in h:
            ws.conditional_formatting.add(h["range"], CellIsRule(operator="lessThan", formula=[str(h["below"])], fill=fill))
    for ch in spec.get("charts", []):
        add_chart(ws, ch)


def add_chart(ws, ch):
    kind = ch.get("type", "column")
    min_col, min_row, max_col, max_row = range_boundaries(ch["data"])
    if kind == "pie":
        chart = PieChart()
    elif kind == "line":
        chart = LineChart()
    elif kind == "scatter":
        chart = ScatterChart()
        chart.style = 13
    else:
        chart = BarChart()
        chart.type = "bar" if kind == "bar" else "col"
    chart.title = ch.get("title")
    chart.height, chart.width = ch.get("height", 7.5), ch.get("width", 15)
    if kind == "scatter":
        cats = ch.get("categories")
        if not cats:
            raise SystemExit("scatter charts need `categories` (the x values)")
        c0, r0, c1, r1 = range_boundaries(cats)
        xs = Reference(ws, min_col=c0, min_row=r0, max_col=c1, max_row=r1)
        for col in range(min_col, max_col + 1):
            ys = Reference(ws, min_col=col, min_row=min_row, max_row=max_row)
            chart.series.append(Series(ys, xs, title_from_data=True))
    else:
        data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        if ch.get("categories"):
            c0, r0, c1, r1 = range_boundaries(ch["categories"])
            chart.set_categories(Reference(ws, min_col=c0, min_row=r0, max_col=c1, max_row=r1))
    if ch.get("x_title") and hasattr(chart, "x_axis"):
        chart.x_axis.title = ch["x_title"]
    if ch.get("y_title") and hasattr(chart, "y_axis"):
        chart.y_axis.title = ch["y_title"]
    ws.add_chart(chart, ch.get("anchor", "H2"))


def read_input(path):
    """The input file's text; `-` reads standard input."""
    if path == "-":
        return sys.stdin.read()
    with open(path, encoding="utf-8") as f:
        return f.read()


def is_json(path, src):
    return path.lower().endswith(".json") or (path == "-" and src.lstrip().startswith("{"))


def cmd_create(a):
    spec = json.loads(read_input(a.input))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in spec.get("sheets", []):
        build_sheet(wb.create_sheet(s.get("name", f"Sheet{len(wb.sheetnames) + 1}")[:31]), s)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    recalc_on_open(wb)
    wb.save(a.output)
    print(f"wrote {a.output} ({', '.join(wb.sheetnames)})")


def cmd_add_sheet(a):
    spec = json.loads(read_input(a.input))
    wb = openpyxl.load_workbook(a.file)
    for s in spec.get("sheets", []):
        build_sheet(wb.create_sheet(s.get("name", f"Sheet{len(wb.sheetnames) + 1}")[:31]), s)
    recalc_on_open(wb)
    wb.save(a.output or a.file)
    print(f"sheets: {', '.join(wb.sheetnames)}")


def cmd_from_csv(a):
    with open(a.input, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if not rows:
        raise SystemExit("empty CSV")
    spec = {
        "name": a.sheet,
        "columns": [{"header": h, "width": min(max(len(h) + 4, 10), 40)} for h in rows[0]],
        "rows": [[parse_scalar(v) if v != "" else None for v in r] for r in rows[1:]],
        "freeze": "A2",
        "filter": True,
    }
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_sheet(wb.create_sheet(a.sheet[:31]), spec)
    wb.save(a.output)
    print(f"wrote {a.output} ({len(rows) - 1} rows)")


def cmd_dump(a):
    wb = openpyxl.load_workbook(a.file)
    out = {}
    for ws in wb.worksheets:
        if a.sheet and ws.title != a.sheet:
            continue
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    v = c.value
                    if isinstance(v, (dt.date, dt.datetime)):
                        v = v.isoformat()
                    cells.append([c.coordinate, v])
        out[ws.title] = {"dimensions": ws.dimensions, "cells": cells, "charts": len(ws._charts)}
    if a.json:
        print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
        return
    for name, s in out.items():
        print(f"--- {name} ({s['dimensions']}, {s['charts']} chart(s))")
        for coord, v in s["cells"]:
            print(f"  {coord}: {v}")


def split_target(wb, target):
    if "!" in target:
        sheet, coord = target.rsplit("!", 1)
        sheet = sheet.strip("'")
        if sheet not in wb.sheetnames:
            raise SystemExit(f"no sheet named {sheet!r} (have {', '.join(wb.sheetnames)})")
        return wb[sheet], coord
    return wb.worksheets[0], target


def cmd_set(a):
    wb = openpyxl.load_workbook(a.file)
    for assignment in a.assignments:
        if "=" not in assignment:
            raise SystemExit(f"expected Sheet!A1=value, got {assignment!r}")
        target, value = assignment.split("=", 1)
        ws, coord = split_target(wb, target)
        ws[coord] = parse_scalar(value)
        if isinstance(ws[coord].value, dt.date) and ws[coord].number_format == "General":
            ws[coord].number_format = "yyyy-mm-dd"
    recalc_on_open(wb)
    wb.save(a.output or a.file)
    print(f"set {len(a.assignments)} cell(s)")


def cmd_append(a):
    wb = openpyxl.load_workbook(a.file)
    if a.sheet not in wb.sheetnames:
        raise SystemExit(f"no sheet named {a.sheet!r}")
    ws = wb[a.sheet]
    rows = json.loads(a.row)
    if rows and not isinstance(rows[0], list):
        rows = [rows]
    for row in rows:
        ws.append([coerce(v) for v in row])
        r = ws.max_row
        for j in range(1, len(row) + 1):
            above = ws.cell(row=r - 1, column=j)
            if r > 2 and above.number_format != "General":
                ws.cell(row=r, column=j).number_format = above.number_format
    recalc_on_open(wb)
    wb.save(a.output or a.file)
    print(f"appended {len(rows)} row(s)")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("create", help="JSON spec → .xlsx")
    c.add_argument("input")
    c.add_argument("output")
    c.set_defaults(fn=cmd_create)
    s = sub.add_parser("add-sheet", help="add a spec's sheets to a workbook")
    s.add_argument("file")
    s.add_argument("input")
    s.add_argument("-o", "--output")
    s.set_defaults(fn=cmd_add_sheet)
    f = sub.add_parser("from-csv", help="CSV → styled .xlsx")
    f.add_argument("input")
    f.add_argument("output")
    f.add_argument("--sheet", default="Data")
    f.set_defaults(fn=cmd_from_csv)
    d = sub.add_parser("dump", help="print values and formulas")
    d.add_argument("file")
    d.add_argument("--sheet")
    d.add_argument("--json", action="store_true")
    d.set_defaults(fn=cmd_dump)
    t = sub.add_parser("set", help="set cells: Sheet!A1=value ...")
    t.add_argument("file")
    t.add_argument("assignments", nargs="+")
    t.add_argument("-o", "--output")
    t.set_defaults(fn=cmd_set)
    p = sub.add_parser("append", help="append a JSON row (or list of rows)")
    p.add_argument("file")
    p.add_argument("sheet")
    p.add_argument("row")
    p.add_argument("-o", "--output")
    p.set_defaults(fn=cmd_append)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
